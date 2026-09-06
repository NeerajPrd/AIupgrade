from typing import List
from loguru import logger
from docx import Document as DocxReader
from starlette.concurrency import run_in_threadpool

from src.utils.common import get_bytes_stream
from src.schemas.document import PageData


class DOCXExtractionError(Exception):
    """Custom exception for DOCX extraction-related failures."""


class DOCXDocumentExtractor:
    """Extracts text from a Word (.docx) file's paragraphs and tables."""

    def __init__(self, file_bytes: bytes) -> None:
        if not file_bytes:
            logger.error("Initialization failed: file_bytes is empty.")
            raise ValueError("file_bytes must not be empty")

        self._file_bytes = file_bytes
        logger.debug("DOCXDocumentExtractor initialized.")

    async def extract_pages(self) -> List[PageData]:
        """Asynchronously extracts text from the file as a single page."""
        return await run_in_threadpool(self._extract_pages_sync)

    def _extract_pages_sync(self) -> List[PageData]:
        try:
            with get_bytes_stream(self._file_bytes) as docx_stream:
                document = DocxReader(docx_stream)
                logger.info("DOCX stream opened successfully.")

                paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

                table_lines = []
                for table in document.tables:
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        if any(cells):
                            table_lines.append(" | ".join(cells))

                text = "\n".join(paragraphs + table_lines)

                return [
                    PageData(
                        content=text,
                        metadata={
                            "paragraph_count": len(paragraphs),
                            "table_count": len(document.tables),
                        },
                    )
                ]

        except Exception as e:
            logger.exception("Unexpected error during DOCX extraction.")
            raise DOCXExtractionError(f"Unexpected error: {e}") from e
