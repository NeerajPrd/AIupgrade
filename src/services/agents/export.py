"""Build a Word document report of an agent's conversations for a date range."""

import io
from datetime import datetime
from typing import Any, Dict, List

from docx import Document


def build_conversations_docx(agent_name: str, period_label: str, conversations: List[Dict[str, Any]]) -> io.BytesIO:
    """
    conversations: [{"title", "created_at" (iso str), "messages": [{"role", "content", "created_at"}]}]
    """
    document = Document()

    document.add_heading(f"{agent_name} — Conversations ({period_label})", level=0)
    document.add_paragraph(
        f"Exported {datetime.now().strftime('%Y-%m-%d %H:%M')} · {len(conversations)} conversation(s)"
    )

    if not conversations:
        document.add_paragraph("No conversations in this period.")

    for convo in conversations:
        document.add_heading(convo.get("title") or "Untitled conversation", level=1)

        started = _format_timestamp(convo.get("created_at"))
        if started:
            meta = document.add_paragraph()
            meta.add_run(f"Started: {started}").italic = True

        for msg in convo.get("messages", []):
            role = msg.get("role")
            label = "Client" if role == "user" else agent_name if role == "assistant" else (role or "Unknown").title()
            timestamp = _format_timestamp(msg.get("created_at"))

            p = document.add_paragraph()
            p.add_run(f"{label} ({timestamp}): ").bold = True
            p.add_run(msg.get("content") or "")

        document.add_paragraph("")

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def _format_timestamp(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return value
    return value.strftime("%Y-%m-%d %H:%M")
